"""Renders source documents to the file formats a real company would store.

Output is byte-for-byte reproducible (fixed PDF metadata, fixed DOCX zip timestamps), so
re-running the generator only re-uploads objects whose content or ACL actually changed.
"""

import io
import json
import zipfile
from collections.abc import Callable
from html import escape as html_escape
from xml.sax.saxutils import escape as xml_escape

from docx import Document
from reportlab.lib import colors
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Flowable, Paragraph, SimpleDocTemplate, Spacer, TableStyle
from reportlab.platypus import Table as PdfTable

from data_gen.models import SourceDocument, Table

_ZIP_TIMESTAMP = (2026, 1, 1, 0, 0, 0)


def render(doc: SourceDocument) -> bytes:
    renderers: dict[str, Callable[[SourceDocument], bytes]] = {
        "pdf": render_pdf,
        "docx": render_docx,
        "md": render_markdown,
        "html": render_html,
        "json": render_json,
    }
    return renderers[doc.key.rsplit(".", 1)[-1]](doc)


def _meta_line(doc: SourceDocument) -> str:
    return (
        f"Owner: {doc.owner} | Last updated: {doc.updated_at:%B %d, %Y} | "
        f"Classification: {doc.classification}"
    )


def render_pdf(doc: SourceDocument) -> bytes:
    styles = getSampleStyleSheet()
    story: list[Flowable] = [
        Paragraph(xml_escape(doc.title), styles["Title"]),
        Paragraph(xml_escape(_meta_line(doc)), styles["Italic"]),
        Spacer(1, 12),
    ]
    for sec in doc.sections:
        story.append(Paragraph(xml_escape(sec.heading), styles["Heading2"]))
        for block in sec.blocks:
            if isinstance(block, Table):
                table = PdfTable([list(block.header), *map(list, block.rows)], repeatRows=1)
                table.setStyle(
                    TableStyle(
                        [
                            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                            ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                        ]
                    )
                )
                story += [table, Spacer(1, 8)]
            else:
                story.append(Paragraph(xml_escape(block), styles["BodyText"]))
    buffer = io.BytesIO()
    template = SimpleDocTemplate(
        buffer, pagesize=LETTER, title=doc.title, author=doc.owner, invariant=1
    )
    template.build(story)
    return buffer.getvalue()


def render_docx(doc: SourceDocument) -> bytes:
    document = Document()
    props = document.core_properties
    props.title = doc.title
    props.author = doc.owner
    props.last_modified_by = doc.owner
    props.created = doc.updated_at.replace(tzinfo=None)
    props.modified = doc.updated_at.replace(tzinfo=None)
    props.revision = 1
    document.add_heading(doc.title, level=0)
    document.add_paragraph().add_run(_meta_line(doc)).italic = True
    for sec in doc.sections:
        document.add_heading(sec.heading, level=1)
        for block in sec.blocks:
            if isinstance(block, Table):
                table = document.add_table(rows=1, cols=len(block.header))
                table.style = "Table Grid"
                for cell, text in zip(table.rows[0].cells, block.header, strict=True):
                    cell.text = text
                for row in block.rows:
                    for cell, text in zip(table.add_row().cells, row, strict=True):
                        cell.text = text
            else:
                document.add_paragraph(block)
    buffer = io.BytesIO()
    document.save(buffer)
    return _normalize_zip(buffer.getvalue())


def _normalize_zip(data: bytes) -> bytes:
    """Rewrites zip entries with a fixed timestamp; python-docx stamps the current time."""
    out = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(data)) as src,
        zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst,
    ):
        for info in src.infolist():
            entry = zipfile.ZipInfo(info.filename, date_time=_ZIP_TIMESTAMP)
            entry.compress_type = zipfile.ZIP_DEFLATED
            entry.external_attr = info.external_attr
            dst.writestr(entry, src.read(info.filename))
    return out.getvalue()


def _md_cell(text: str) -> str:
    return text.replace("|", "\\|")


def render_markdown(doc: SourceDocument) -> bytes:
    lines = [
        "---",
        f"title: {json.dumps(doc.title)}",
        f"owner: {doc.owner}",
        f"updated: {doc.updated_at:%Y-%m-%d}",
        "---",
        "",
        f"# {doc.title}",
        "",
    ]
    for sec in doc.sections:
        lines += [f"## {sec.heading}", ""]
        for block in sec.blocks:
            if isinstance(block, Table):
                lines.append("| " + " | ".join(map(_md_cell, block.header)) + " |")
                lines.append("|" + "---|" * len(block.header))
                lines += ["| " + " | ".join(map(_md_cell, row)) + " |" for row in block.rows]
            else:
                lines.append(block)
            lines.append("")
    return "\n".join(lines).encode()


def render_html(doc: SourceDocument) -> bytes:
    """A legacy wiki export, including the navigation and footer chrome parsers must strip."""
    e = html_escape
    parts = [
        "<!DOCTYPE html>",
        '<html><head><meta charset="utf-8">',
        f"<title>{e(doc.title)}</title></head><body>",
        f'<div id="header"><nav>Larkspur Wiki &raquo; Spaces &raquo; {e(doc.owner)}</nav></div>',
        '<div id="main-content" class="wiki-content">',
        f"<h1>{e(doc.title)}</h1>",
        f'<p class="page-metadata">Last updated {doc.updated_at:%Y-%m-%d} by {e(doc.owner)}</p>',
    ]
    for sec in doc.sections:
        parts.append(f"<h2>{e(sec.heading)}</h2>")
        for block in sec.blocks:
            if isinstance(block, Table):
                parts.append("<table>")
                parts.append("<tr>" + "".join(f"<th>{e(h)}</th>" for h in block.header) + "</tr>")
                parts += [
                    "<tr>" + "".join(f"<td>{e(c)}</td>" for c in row) + "</tr>"
                    for row in block.rows
                ]
                parts.append("</table>")
            else:
                parts.append(f"<p>{e(block)}</p>")
    parts += [
        "</div>",
        '<div id="footer">Exported from Larkspur Wiki. Powered by Confluence.</div>',
        "</body></html>",
    ]
    return "\n".join(parts).encode()


def render_json(doc: SourceDocument) -> bytes:
    return json.dumps(doc.payload, indent=2, ensure_ascii=False).encode()
